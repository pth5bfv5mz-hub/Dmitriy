#!/usr/bin/env python3
"""Служебный просмотр вёрстки презентации без PowerPoint.

Рисует каждый слайд собранной колоды в PNG по геометрии фигур из .pptx
и предупреждает о выходе текста и фигур за границы слайда.

    python scripts/preview_deck.py --lesson 3
    python scripts/preview_deck.py --lesson 3 --outdir /tmp/preview

Инструмент разработчика: пишет только в указанную папку (по умолчанию —
во временную), в output/ и site/ не пишет. Точность передачи шрифтов
приблизительная, назначение — контроль переполнения и композиции.
"""
from __future__ import annotations

import argparse
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

from lib import loader, paths

SCALE = 110  # пикселей на дюйм
FONT_REG = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
FONT_MONO = "/usr/share/fonts/truetype/liberation/LiberationMono-Bold.ttf"

_font_cache: dict = {}


def font(size_pt: float, *, bold=False, mono=False):
    px = max(7, int(size_pt * SCALE / 72))
    key = (px, bold, mono)
    if key not in _font_cache:
        path = FONT_MONO if mono else (FONT_BOLD if bold else FONT_REG)
        _font_cache[key] = ImageFont.truetype(path, px)
    return _font_cache[key]


def px(emu) -> int:
    return int(Emu(int(emu)).inches * SCALE)


def wrap(draw, text, fnt, max_w):
    words, lines, cur = str(text).split(), [], ""
    for w in words:
        probe = f"{cur} {w}".strip()
        if draw.textlength(probe, font=fnt) <= max_w or not cur:
            cur = probe
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def draw_table(draw, shape, warn, slide_no):
    table = shape.table
    x0, y0 = px(shape.left), px(shape.top)
    col_w = [px(c.width) for c in table.columns]
    row_h = [px(r.height) for r in table.rows]
    y = y0
    for r, row in enumerate(table.rows):
        x = x0
        for c, cell in enumerate(row.cells):
            w, h = col_w[c], row_h[r]
            fill = None
            try:
                if cell.fill.type is not None and cell.fill.type == 1:
                    fill = "#" + str(cell.fill.fore_color.rgb)
            except Exception:
                fill = None
            draw.rectangle([x, y, x + w, y + h], fill=fill, outline="#B9C4D4")
            runs = [run for p in cell.text_frame.paragraphs for run in p.runs]
            if runs:
                run = runs[0]
                size = run.font.size.pt if run.font.size else 18
                fnt = font(size, bold=bool(run.font.bold))
                color = "#1A1A1A"
                try:
                    if run.font.color and run.font.color.rgb:
                        color = "#" + str(run.font.color.rgb)
                except Exception:
                    pass
                lines = wrap(draw, cell.text, fnt, w - 14)
                lh = int(size * SCALE / 72 * 1.22)
                ty = y + max(3, (h - lh * len(lines)) // 2)
                for line in lines:
                    draw.text((x + 7, ty), line, font=fnt, fill=color)
                    ty += lh
                if lh * len(lines) > h:
                    warn.append(f"слайд {slide_no}: текст ячейки таблицы "
                                f"не вмещается по высоте: «{cell.text[:40]}»")
            x += w
        y += row_h[r]


def draw_textframe(draw, shape, warn, slide_no, canvas_h):
    x, y = px(shape.left), px(shape.top)
    w, h = px(shape.width), px(shape.height)
    cursor = y
    for para in shape.text_frame.paragraphs:
        runs = para.runs
        if not runs:
            continue
        run = runs[0]
        size = run.font.size.pt if run.font.size else 18
        mono = bool(run.font.name and "Consolas" in str(run.font.name))
        fnt = font(size, bold=bool(run.font.bold), mono=mono)
        color = "#1A1A1A"
        try:
            if run.font.color and run.font.color.rgb:
                color = "#" + str(run.font.color.rgb)
        except Exception:
            pass
        text = "".join(r.text for r in runs)
        before = para.space_before.pt if para.space_before else 0
        after = para.space_after.pt if para.space_after else 0
        cursor += int(before * SCALE / 72)
        lines = wrap(draw, text, fnt, w - 4)
        lh = int(size * SCALE / 72 * 1.22)
        align = str(para.alignment or "")
        for line in lines:
            tw = draw.textlength(line, font=fnt)
            tx = x
            if "CENTER" in align:
                tx = x + (w - tw) / 2
            elif "RIGHT" in align:
                tx = x + w - tw
            draw.text((tx, cursor), line, font=fnt, fill=color)
            cursor += lh
        cursor += int(after * SCALE / 72)
    if cursor > canvas_h - 4:
        warn.append(f"слайд {slide_no}: текст выходит за нижний край слайда "
                    f"({cursor - canvas_h} px)")
    return cursor


def render(pptx_path: Path, outdir: Path) -> list[str]:
    prs = Presentation(str(pptx_path))
    W, H = px(prs.slide_width), px(prs.slide_height)
    warn: list[str] = []
    outdir.mkdir(parents=True, exist_ok=True)
    made = []

    for i, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (W, H), "white")
        draw = ImageDraw.Draw(img)
        draw.rectangle([0, 0, W - 1, H - 1], outline="#D0D7E2")

        for shape in slide.shapes:
            if shape.has_table:
                draw_table(draw, shape, warn, i)
                continue
            if shape.shape_type is not None and not shape.has_text_frame:
                continue
            is_autoshape = shape.shape_type is not None and shape.shape_type != 17
            if is_autoshape:
                fill = None
                try:
                    if shape.fill.type == 1:
                        fill = "#" + str(shape.fill.fore_color.rgb)
                except Exception:
                    pass
                draw.rectangle([px(shape.left), px(shape.top),
                                px(shape.left) + px(shape.width),
                                px(shape.top) + px(shape.height)], fill=fill)
            if shape.has_text_frame and shape.text_frame.text.strip():
                draw_textframe(draw, shape, warn, i, H)
            r = px(shape.left) + px(shape.width)
            b = px(shape.top) + px(shape.height)
            if r > W + 2 or b > H + 2:
                warn.append(f"слайд {i}: фигура выходит за границы слайда "
                            f"(право {r - W} px, низ {b - H} px)")

        path = outdir / f"slide_{i:02d}.png"
        img.save(path)
        made.append(path)

    contact = outdir / "all_slides.png"
    cols = 3
    rows = (len(made) + cols - 1) // cols
    thumb_w, thumb_h = W // 2, H // 2
    sheet = Image.new("RGB", (cols * thumb_w, rows * thumb_h), "#E9EDF3")
    for idx, p in enumerate(made):
        im = Image.open(p).resize((thumb_w, thumb_h))
        sheet.paste(im, ((idx % cols) * thumb_w, (idx // cols) * thumb_h))
    sheet.save(contact)
    return warn


def main() -> int:
    ap = argparse.ArgumentParser(description="Просмотр вёрстки презентации")
    ap.add_argument("--lesson", type=int, required=True)
    ap.add_argument("--discipline", default="mdk-04-02")
    ap.add_argument("--outdir")
    args = ap.parse_args()

    discipline = loader.load_discipline(args.discipline)
    deck = (paths.lesson_dir(discipline, args.lesson)
            / f"Занятие_{args.lesson:02d}_Презентация.pptx")
    if not deck.exists():
        raise SystemExit(f"Нет файла: {deck}")

    outdir = Path(args.outdir) if args.outdir else Path(
        tempfile.mkdtemp(prefix=f"deck{args.lesson:02d}_"))
    warns = render(deck, outdir)
    print(f"Картинки слайдов: {outdir}")
    if warns:
        print(f"\nПредупреждения ({len(warns)}):")
        for w in dict.fromkeys(warns):
            print(f"  {w}")
    else:
        print("Переполнения и выходов за границы не обнаружено.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
