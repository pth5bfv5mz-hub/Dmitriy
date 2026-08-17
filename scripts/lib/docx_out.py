"""Генераторы .docx: технологическая карта и практическая часть занятия."""
from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import paths

CONDITIONAL_NOTE = "Все числовые данные условные, используются в учебных целях."


# --- низкоуровневые помощники ---------------------------------------------

def _shade(cell, hex_color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _cell_text(cell, text, *, bold=False, size=None, color=None,
               align=WD_ALIGN_PARAGRAPH.LEFT, font=None, italic=False):
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = align
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.space_before = Pt(2)
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        if i:
            para = cell.add_paragraph()
            para.alignment = align
            para.paragraph_format.space_after = Pt(2)
        run = para.add_run(line)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if font:
            run.font.name = font


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))


class DocBuilder:
    """Обёртка над python-docx с оформлением курса."""

    def __init__(self, theme: dict):
        self.theme = theme
        self.p = theme["palette"]
        self.cfg = theme["docx"]
        self.doc = Document(str(paths.DOC_TEMPLATE))
        self._content_width_cm = 21.0 - (self.cfg["page"]["margins_cm"]["left"]
                                         + self.cfg["page"]["margins_cm"]["right"])

    # --- текст ---

    def title(self, text: str, subtitle: str | None = None):
        para = self.doc.add_paragraph(style="Курс Титул")
        para.add_run(text)
        if subtitle:
            sub = self.doc.add_paragraph(style="Курс Подзаголовок")
            sub.add_run(subtitle)

    def heading(self, text: str, level: int = 1):
        self.doc.add_heading(text, level=level)

    def para(self, text: str = "", *, bold=False, italic=False, size=None,
             style=None, align=None, space_after=None):
        para = self.doc.add_paragraph(style=style)
        if align is not None:
            para.alignment = align
        if space_after is not None:
            para.paragraph_format.space_after = Pt(space_after)
        if text:
            run = para.add_run(text)
            run.bold = bold
            run.italic = italic
            if size:
                run.font.size = Pt(size)
        return para

    def caption(self, text: str):
        return self.para(text, style="Курс Подпись", italic=True)

    def bullets(self, items, *, numbered=False, size=None):
        style = "List Number" if numbered else "List Bullet"
        for item in items:
            para = self.doc.add_paragraph(style=style)
            para.paragraph_format.space_after = Pt(3)
            run = para.add_run(str(item))
            if size:
                run.font.size = Pt(size)

    def formula(self, text: str, legend=None, *, units: str | None = None):
        para = self.doc.add_paragraph(style="Курс Формула")
        run = para.add_run(text)
        run.font.name = "Consolas"
        if legend:
            self.para("где:", italic=True, size=self.cfg["font_size_small"],
                      space_after=2)
            for item in legend:
                p = self.doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1.0)
                p.paragraph_format.space_after = Pt(1)
                r = p.add_run(str(item))
                r.font.size = Pt(self.cfg["font_size_small"])
        if units:
            self.caption(f"Единица измерения: {units}")

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def rule(self):
        para = self.doc.add_paragraph()
        para.paragraph_format.space_after = Pt(6)
        p_pr = para._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), self.p["accent"])
        borders.append(bottom)
        p_pr.append(borders)

    # --- таблицы ---

    def table(self, header, rows, *, widths=None, header_fill=None,
              zebra=True, size=None, first_col_bold=False, align=None):
        n_cols = len(header) if header else len(rows[0])
        table = self.doc.add_table(rows=0, cols=n_cols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        size = size or self.cfg["font_size_small"]

        if widths:
            total = sum(widths)
            col_cm = [self._content_width_cm * w / total for w in widths]
            for i, cm in enumerate(col_cm):
                for cell in table.columns[i].cells:
                    cell.width = Cm(cm)

        if header:
            row = table.add_row()
            _repeat_header(row)
            for i, text in enumerate(header):
                _cell_text(row.cells[i], text, bold=True,
                           color=self.p["white"], size=size,
                           align=WD_ALIGN_PARAGRAPH.CENTER)
                _shade(row.cells[i], header_fill or self.p["primary"])

        aligns = align or []
        for r, data_row in enumerate(rows):
            row = table.add_row()
            for c, text in enumerate(data_row):
                al = WD_ALIGN_PARAGRAPH.LEFT
                if c < len(aligns):
                    al = {"l": WD_ALIGN_PARAGRAPH.LEFT,
                          "c": WD_ALIGN_PARAGRAPH.CENTER,
                          "r": WD_ALIGN_PARAGRAPH.RIGHT}.get(
                              aligns[c], WD_ALIGN_PARAGRAPH.LEFT)
                _cell_text(row.cells[c], text, size=size, align=al,
                           bold=(first_col_bold and c == 0))
                if zebra and r % 2 == 1:
                    _shade(row.cells[c], self.p["surface"])
            if widths:
                col_cm = [self._content_width_cm * w / sum(widths) for w in widths]
                for i, cm in enumerate(col_cm):
                    row.cells[i].width = Cm(cm)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return table

    def kv_table(self, pairs, *, widths=(30, 70)):
        rows = [[k, v] for k, v in pairs]
        return self.table(None, rows, widths=list(widths), zebra=False,
                          first_col_bold=True, size=self.cfg["font_size"])

    # --- колонтитул и сохранение ---

    def footer(self, text: str):
        section = self.doc.sections[0]
        para = section.footer.paragraphs[0]
        para.text = ""
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(text)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(self.p["muted"])

    def save(self, path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        return path


# --- Технологическая карта -------------------------------------------------

_BLOCK_TITLES = {
    "org": "Организационная часть",
    "theory": "Блок 1. Теория (изложение по презентации)",
    "practice": "Блок 2. Практика (задание по теории занятия)",
    "final": "Заключительная часть",
}


def _competency_rows(cfg, competencies):
    rows = []
    unresolved = []
    for kind, registry_key in (("pk", "pk"), ("ok", "ok")):
        for code in (cfg.get("competencies") or {}).get(kind, []) or []:
            entry = (competencies.get(registry_key) or {}).get(code)
            if not entry:
                placeholder = "[ПК-?]" if kind == "pk" else "[ОК-?]"
                rows.append([placeholder, f"код «{code}» не найден в реестре"])
                unresolved.append(code)
                continue
            text = entry["text"]
            if entry.get("verified") == "partial":
                text += "  [сверить формулировку с РП]"
            elif entry.get("verified") == "none":
                code = "[ПК-?]" if kind == "pk" else "[ОК-?]"
            rows.append([code, text])
    return rows, unresolved


def build_tech_card(theme, discipline, row, cfg, competencies, out_path):
    """Техкарта занятия. Возвращает (путь, список непроставленных кодов)."""
    b = DocBuilder(theme)
    n = row.n
    b.footer(f'{discipline.code} «{discipline.title}» · Занятие {n} · '
             f'{discipline.group}')

    b.title("Технологическая карта занятия",
            f'{discipline.code} «{discipline.title}»')

    b.kv_table([
        ("Специальность", f'{discipline.specialty_code} {discipline.specialty}'),
        ("Профессиональный модуль", f'{discipline.module}'),
        ("Группа, курс", f'{discipline.group}, {discipline.course} курс, '
                         f'{discipline.form} форма'),
        ("Раздел", row.section),
        ("Занятие №", str(n)),
        ("Тема занятия", row.topic),
        ("Количество часов", f'{row.hours} ак. ч ({45 * row.hours} мин)'),
        ("Тип занятия", cfg["lesson_type"]),
        ("Форма организации", cfg.get("lesson_form",
                                      "фронтальная, индивидуальная")),
        ("Методы обучения", cfg.get("methods",
                                    "объяснительно-иллюстративный, "
                                    "репродуктивный, практический")),
        ("Контрольная точка", row.checkpoint or "—"),
    ])

    b.heading("1. Цели занятия", 1)
    goals = cfg["goals"]
    b.kv_table([
        ("Обучающая", goals["teaching"]),
        ("Развивающая", goals["developing"]),
        ("Воспитательная", goals["upbringing"]),
    ], widths=(24, 76))

    b.heading("2. Формируемые компетенции", 1)
    comp_rows, unresolved = _competency_rows(cfg, competencies)
    if comp_rows:
        b.table(["Код", "Формулировка"], comp_rows, widths=[13, 87])
    else:
        b.para("Компетенции не заданы в конфигурации занятия.", italic=True)
    std = competencies.get("standard") or {}
    if std:
        b.caption(f'Формулировки по ФГОС СПО {std.get("code", "")} '
                  f'{std.get("name", "")} ({std.get("order", "")}).')

    b.heading("3. Оснащение занятия", 1)
    b.bullets(cfg["equipment"])

    b.heading("4. Планируемые результаты", 1)
    results = cfg.get("results") or {}
    if results:
        b.kv_table([
            ("Знать", results.get("know", "—")),
            ("Уметь", results.get("can", "—")),
        ], widths=(20, 80))

    b.heading("5. Ход занятия", 1)
    b.caption(f'Структура пары: теория 40–50 мин, далее практика по теории '
              f'этого же занятия. Всего {45 * row.hours} мин.')

    stages = cfg["stages"]
    current_block = None
    for stage in stages:
        block = stage.get("block")
        if block != current_block:
            current_block = block
            b.heading(_BLOCK_TITLES.get(block, block), 2)
            rows_buf = []
            for st in stages:
                if st.get("block") != block:
                    continue
                rows_buf.append([
                    f'{st["stage"]}\n({st["time"]})',
                    st["teacher"],
                    st["student"],
                ])
            b.table(["Этап и время", "Деятельность преподавателя",
                     "Деятельность студентов"],
                    rows_buf, widths=[20, 40, 40])

    b.heading("6. Домашнее задание", 1)
    hw = cfg["homework"]
    if isinstance(hw, dict):
        b.para(hw.get("text", ""))
        if hw.get("bullets"):
            b.bullets(hw["bullets"])
        if hw.get("source"):
            b.caption(hw["source"])
    else:
        b.para(str(hw))

    b.heading("7. Литература и источники", 1)
    lit = cfg.get("literature") or []
    if lit:
        b.bullets(lit, numbered=True, size=theme["docx"]["font_size_small"])
    else:
        b.para("—")

    notes = cfg.get("teacher_notes")
    if notes:
        b.heading("8. Примечания преподавателю", 1)
        b.bullets(notes, size=theme["docx"]["font_size_small"])

    return b.save(out_path), unresolved


# --- Практическая часть ----------------------------------------------------

def _practice_header(b, theme, discipline, row, cfg, practice, competencies):
    kind_title = {
        "task": "Практическое задание (расчётная задача)",
        "case": "Практическое задание (разбор кейса)",
        "test": "Тест по теме занятия",
    }[practice["kind"]]
    b.footer(f'{discipline.code} «{discipline.title}» · Занятие {row.n} · '
             f'{discipline.group}')
    b.title(kind_title, f'{discipline.code} «{discipline.title}»')
    b.kv_table([
        ("Занятие №", str(row.n)),
        ("Тема", row.topic),
        ("Группа", discipline.group),
        ("Время на выполнение", practice.get("duration", "40 мин")),
        ("Форма работы", practice.get("form", "индивидуально")),
    ])
    b.heading("Цель работы", 1)
    b.para(practice["goal"])
    comp_rows, unresolved = _competency_rows(cfg, competencies)
    if comp_rows:
        b.heading("Проверяемые компетенции", 1)
        b.para(", ".join(r[0] for r in comp_rows))
    return unresolved


def _render_data_blocks(b, blocks):
    for block in blocks or []:
        if block.get("title"):
            b.para(block["title"], bold=True, space_after=3)
        if block.get("text"):
            b.para(block["text"])
        if block.get("table"):
            t = block["table"]
            b.table(t.get("header"), t["rows"], widths=t.get("widths"),
                    align=t.get("align"))
        if block.get("note"):
            b.caption(block["note"])
        if block.get("bullets"):
            b.bullets(block["bullets"])


def build_practice(theme, discipline, row, cfg, competencies, out_path):
    practice = cfg["practice"]
    b = DocBuilder(theme)
    unresolved = _practice_header(b, theme, discipline, row, cfg, practice,
                                 competencies)

    if practice["kind"] == "test":
        _build_test_body(b, theme, practice)
    else:
        _build_task_body(b, theme, practice)

    return b.save(out_path), unresolved


def _build_task_body(b, theme, practice):
    b.heading("Условие", 1)
    b.para(CONDITIONAL_NOTE, italic=True,
           size=theme["docx"]["font_size_small"])
    if practice.get("situation"):
        b.para(practice["situation"])
    _render_data_blocks(b, practice.get("data"))

    if practice.get("questions"):
        b.heading("Что требуется определить", 1)
        b.bullets(practice["questions"], numbered=True)

    b.heading("Порядок выполнения", 1)
    b.bullets(practice["steps"], numbered=True)

    if practice.get("formulas"):
        b.heading("Формулы для расчёта", 1)
        for item in practice["formulas"]:
            if item.get("title"):
                b.para(item["title"], bold=True, space_after=2)
            b.formula(item["formula"], item.get("legend"),
                      units=item.get("units"))

    b.heading("Критерии оценки", 1)
    b.table(["Оценка", "Критерий"], practice["criteria"], widths=[18, 82],
            align=["c", "l"])
    if practice.get("criteria_note"):
        b.caption(practice["criteria_note"])

    if practice.get("common_mistakes"):
        b.heading("Типичные ошибки", 1)
        b.bullets(practice["common_mistakes"],
                  size=theme["docx"]["font_size_small"])

    # Эталон — отдельным разделом в конце, с новой страницы
    b.page_break()
    b.heading("Эталон решения", 1)
    b.caption("Раздел для преподавателя. При выдаче студентам страницу "
              "не печатать.")
    solution = practice["solution"]
    for step in solution.get("steps", []):
        b.para(step["label"], bold=True, space_after=2)
        if step.get("calc"):
            b.formula(step["calc"])
        if step.get("comment"):
            b.para(step["comment"], italic=True,
                   size=theme["docx"]["font_size_small"])
    if solution.get("table"):
        t = solution["table"]
        b.table(t.get("header"), t["rows"], widths=t.get("widths"),
                align=t.get("align"))
    if solution.get("answer"):
        b.rule()
        para = b.doc.add_paragraph(style="Курс Ответ")
        para.add_run(f'Ответ: {solution["answer"]}')
    if solution.get("conclusion"):
        b.para("Вывод для проверки формулировки студента:", bold=True,
               space_after=2)
        b.para(solution["conclusion"])


def _build_test_body(b, theme, practice):
    b.heading("Инструкция", 1)
    b.para(practice.get("instruction",
                        "Вопросы 1–8 — выберите один верный вариант ответа. "
                        "Вопросы 9–10 — дайте краткий письменный ответ."))
    b.para(CONDITIONAL_NOTE, italic=True,
           size=theme["docx"]["font_size_small"])

    b.heading("Вопросы", 1)
    letters = "абвгд"
    for i, q in enumerate(practice["questions"], 1):
        b.para(f'{i}. {q["text"]}', bold=True, space_after=2)
        options = q.get("options")
        if options:
            for j, opt in enumerate(options):
                para = b.doc.add_paragraph()
                para.paragraph_format.left_indent = Cm(1.0)
                para.paragraph_format.space_after = Pt(1)
                para.add_run(f'{letters[j]}) {opt}')
        else:
            para = b.doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(1.0)
            para.paragraph_format.space_after = Pt(10)
            para.add_run("Ответ: " + "_" * 60)

    b.heading("Критерии оценки", 1)
    b.table(["Оценка", "Критерий"], practice["criteria"], widths=[18, 82],
            align=["c", "l"])

    b.page_break()
    b.heading("Ключи к тесту", 1)
    b.caption("Отдельная страница для преподавателя.")
    closed_rows, open_rows = [], []
    for i, q in enumerate(practice["questions"], 1):
        if q.get("options"):
            closed_rows.append([str(i), str(q["answer"])])
        else:
            open_rows.append([str(i), str(q["answer"])])
    if closed_rows:
        b.para("Закрытые вопросы", bold=True, space_after=3)
        b.table(["№", "Верный вариант"], closed_rows, widths=[15, 85],
                align=["c", "l"])
    if open_rows:
        b.para("Открытые вопросы — ожидаемое содержание ответа", bold=True,
               space_after=3)
        b.table(["№", "Ожидаемый ответ"], open_rows, widths=[15, 85],
                align=["c", "l"])


# --- Дополнительные материалы (для занятий без практики) -------------------

def build_materials(theme, discipline, row, cfg, competencies, out_path):
    """Для занятий, где практическая работа не предусмотрена (консультация)."""
    materials = cfg["materials"]
    b = DocBuilder(theme)
    b.footer(f'{discipline.code} «{discipline.title}» · Занятие {row.n} · '
             f'{discipline.group}')
    b.title(materials.get("title", "Материалы к занятию"),
            f'{discipline.code} «{discipline.title}»')
    b.kv_table([
        ("Занятие №", str(row.n)),
        ("Тема", row.topic),
        ("Группа", discipline.group),
    ])
    for section in materials.get("sections", []):
        b.heading(section["heading"], 1)
        if section.get("text"):
            b.para(section["text"])
        if section.get("bullets"):
            b.bullets(section["bullets"], numbered=section.get("numbered", False))
        if section.get("table"):
            t = section["table"]
            b.table(t.get("header"), t["rows"], widths=t.get("widths"),
                    align=t.get("align"))
        if section.get("note"):
            b.caption(section["note"])
    comp_rows, unresolved = _competency_rows(cfg, competencies)
    return b.save(out_path), unresolved
