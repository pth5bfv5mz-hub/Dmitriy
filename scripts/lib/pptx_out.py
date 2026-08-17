"""Генератор презентаций .pptx. Единый шаблон на весь курс.

Типы слайдов (поле `type` в конфиге занятия):
    title      — тема и цели занятия
    question   — слайд-вопрос (в начале)
    bullets    — заголовок и до 5 коротких пунктов
    definition — определение выделенным блоком
    formula    — формула крупно с расшифровкой каждого обозначения
    table      — таблица вместо абзаца
    two_col    — сравнение в две колонки
    scheme     — схема из фигур (цепочка или блоки)
    example    — шаг разобранного числового примера
    check      — «Закрепление»: вопросы
    summary    — итог занятия
    homework   — домашнее задание
"""
from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from . import paths


def _rgb(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str)


class DeckBuilder:
    def __init__(self, theme: dict, discipline, lesson_number: int, topic: str):
        self.theme = theme
        self.p = theme["palette"]
        self.cfg = theme["pptx"]
        self.sizes = self.cfg["sizes"]
        self.font = self.cfg["font"]
        self.discipline = discipline
        self.number = lesson_number
        self.topic = topic

        self.prs = Presentation(str(paths.DECK_TEMPLATE))
        self.W = self.prs.slide_width
        self.H = self.prs.slide_height

        m = self.cfg["margins_in"]
        self.left = Inches(m["left"])
        self.right_edge = self.W - Inches(m["right"])
        self.content_w = self.right_edge - self.left
        self.top = Inches(m["top"])
        self.footer_y = self.H - Inches(m["bottom"])

        self.title_h = Inches(1.05)
        self.body_top = self.top + self.title_h + Inches(0.28)
        self.body_h = self.footer_y - self.body_top - Inches(0.15)

    # --- низкоуровневые помощники -------------------------------------

    def _blank(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        return slide

    def _textbox(self, slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        return tf

    def _para(self, tf, text, *, size, bold=False, color=None, align=PP_ALIGN.LEFT,
              space_after=6, space_before=0, italic=False, first=False,
              font=None, line_spacing=None):
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(space_after)
        para.space_before = Pt(space_before)
        if line_spacing:
            para.line_spacing = line_spacing
        run = para.add_run()
        run.text = str(text)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font or self.font
        run.font.color.rgb = _rgb(color or self.p["text"])
        return para

    def _rect(self, slide, x, y, w, h, fill, *, shape=MSO_SHAPE.RECTANGLE,
              line=None):
        sh = slide.shapes.add_shape(shape, x, y, w, h)
        sh.fill.solid()
        sh.fill.fore_color.rgb = _rgb(fill)
        if line:
            sh.line.color.rgb = _rgb(line)
            sh.line.width = Pt(1)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        sh.text_frame.word_wrap = True
        return sh

    def _title(self, slide, text: str):
        size = (self.sizes["heading"] if len(str(text)) <= 46
                else self.sizes["heading_long"])
        tf = self._textbox(slide, self.left, self.top, self.content_w,
                           self.title_h, anchor=MSO_ANCHOR.BOTTOM)
        self._para(tf, text, size=size, bold=True, color=self.p["primary"],
                   first=True, space_after=0)
        # тонкая акцентная линия под заголовком
        self._rect(slide, self.left, self.top + self.title_h + Inches(0.06),
                   Inches(1.6), Emu(26670), self.p["accent"])

    def _footer(self, slide, index: int):
        text = (f'{self.discipline.code} «{self.discipline.title}»'
                f'  ·  Занятие {self.number}')
        tf = self._textbox(slide, self.left, self.footer_y,
                           self.content_w - Inches(0.8), Inches(0.3))
        self._para(tf, text, size=self.sizes["footer"], color=self.p["muted"],
                   first=True, space_after=0)
        tf2 = self._textbox(slide, self.right_edge - Inches(0.8), self.footer_y,
                            Inches(0.8), Inches(0.3))
        self._para(tf2, str(index), size=self.sizes["footer"],
                   color=self.p["muted"], align=PP_ALIGN.RIGHT, first=True,
                   space_after=0)

    # --- типы слайдов -------------------------------------------------

    def slide_title(self, data):
        slide = self._blank()
        self._rect(slide, 0, 0, self.W, Inches(2.55), self.p["primary"])

        tf = self._textbox(slide, self.left, Inches(0.5),
                           self.content_w, Inches(0.4))
        self._para(tf, f'{self.discipline.code}  ·  Занятие {self.number}',
                   size=16, color=self.p["accent_soft"], first=True,
                   space_after=0)

        tf = self._textbox(slide, self.left, Inches(1.0), self.content_w,
                           Inches(1.4))
        size = self.sizes["title_slide"] if len(self.topic) <= 52 else 32
        self._para(tf, self.topic, size=size, bold=True,
                   color=self.p["white"], first=True, space_after=0)

        y = Inches(3.0)
        tf = self._textbox(slide, self.left, y, self.content_w, Inches(0.45))
        self._para(tf, "Цели занятия", size=self.sizes["subtitle"], bold=True,
                   color=self.p["primary"], first=True, space_after=4)

        y += Inches(0.62)
        goals = data.get("goals") or []
        box_h = self.footer_y - y - Inches(0.35)
        self._rect(slide, self.left, y, Emu(int(Inches(0.055))), box_h,
                   self.p["accent"])
        tf = self._textbox(slide, self.left + Inches(0.28), y + Inches(0.05),
                           self.content_w - Inches(0.4), box_h)
        for i, g in enumerate(goals[:4]):
            self._para(tf, f"•  {g}", size=self.sizes["body"],
                       first=(i == 0), space_after=10)

        tf = self._textbox(slide, self.left, self.footer_y - Inches(0.15),
                           self.content_w, Inches(0.35))
        self._para(tf, f'{self.discipline.group}  ·  {self.discipline.specialty_code} '
                       f'{self.discipline.specialty}',
                   size=14, color=self.p["muted"], first=True, space_after=0)
        return slide

    def slide_question(self, data):
        slide = self._blank()
        self._title(slide, data.get("heading", "Вопрос для начала"))

        y = self.body_top + Inches(0.2)
        h = Inches(1.9)
        self._rect(slide, self.left, y, self.content_w, h, self.p["accent_soft"])
        self._rect(slide, self.left, y, Emu(int(Inches(0.07))), h, self.p["accent"])
        tf = self._textbox(slide, self.left + Inches(0.35), y + Inches(0.3),
                           self.content_w - Inches(0.7), h - Inches(0.5),
                           anchor=MSO_ANCHOR.MIDDLE)
        self._para(tf, data["question"], size=28, bold=True,
                   color=self.p["primary"], first=True, space_after=0)

        hints = data.get("hints") or []
        if hints:
            y2 = y + h + Inches(0.45)
            tf = self._textbox(slide, self.left, y2, self.content_w,
                               self.footer_y - y2 - Inches(0.2))
            self._para(tf, "Подумайте:", size=self.sizes["body_small"],
                       bold=True, color=self.p["muted"], first=True,
                       space_after=8)
            for hint in hints[:4]:
                self._para(tf, f"—  {hint}", size=self.sizes["body"],
                           space_after=8)
        return slide

    def slide_bullets(self, data):
        slide = self._blank()
        self._title(slide, data["heading"])
        y = self.body_top + Inches(0.1)
        tf = self._textbox(slide, self.left, y, self.content_w,
                           self.footer_y - y - Inches(0.2))
        lead = data.get("lead")
        first = True
        if lead:
            self._para(tf, lead, size=self.sizes["body"], color=self.p["muted"],
                       italic=True, first=True, space_after=14)
            first = False
        for b in (data.get("bullets") or [])[:5]:
            self._para(tf, f"•  {b}", size=self.sizes["body"], first=first,
                       space_after=14)
            first = False
        note = data.get("note")
        if note:
            self._para(tf, note, size=self.sizes["body_small"],
                       color=self.p["muted"], italic=True, space_before=10,
                       space_after=0)
        return slide

    def slide_definition(self, data):
        slide = self._blank()
        self._title(slide, data.get("heading", "Определение"))

        y = self.body_top + Inches(0.15)
        term = data.get("term", "")
        text = data["text"]
        h = Inches(2.5)
        self._rect(slide, self.left, y, self.content_w, h, self.p["accent_soft"])
        self._rect(slide, self.left, y, Emu(int(Inches(0.07))), h,
                   self.p["primary"])
        tf = self._textbox(slide, self.left + Inches(0.35), y + Inches(0.28),
                           self.content_w - Inches(0.7), h - Inches(0.5))
        if term:
            self._para(tf, term, size=26, bold=True, color=self.p["primary"],
                       first=True, space_after=10)
            self._para(tf, text, size=self.sizes["body"], space_after=0)
        else:
            self._para(tf, text, size=24, bold=True, color=self.p["primary"],
                       first=True, space_after=0)

        items = data.get("bullets") or []
        if items:
            y2 = y + h + Inches(0.4)
            tf = self._textbox(slide, self.left, y2, self.content_w,
                               self.footer_y - y2 - Inches(0.2))
            for i, b in enumerate(items[:4]):
                self._para(tf, f"•  {b}", size=self.sizes["body"],
                           first=(i == 0), space_after=10)
        return slide

    def slide_formula(self, data):
        slide = self._blank()
        self._title(slide, data.get("heading", "Формула"))

        y = self.body_top + Inches(0.1)
        fh = Inches(1.35)
        self._rect(slide, self.left, y, self.content_w, fh, self.p["formula_bg"])
        self._rect(slide, self.left, y, Emu(int(Inches(0.07))), fh,
                   self.p["formula_bar"])
        tf = self._textbox(slide, self.left + Inches(0.3), y,
                           self.content_w - Inches(0.6), fh,
                           anchor=MSO_ANCHOR.MIDDLE)
        self._para(tf, data["formula"], size=self.sizes["formula"], bold=True,
                   color=self.p["primary"], align=PP_ALIGN.CENTER, first=True,
                   space_after=0, font=self.cfg["font_mono"])

        y += fh + Inches(0.3)
        tf = self._textbox(slide, self.left, y, self.content_w,
                           self.footer_y - y - Inches(0.2))
        self._para(tf, "где:", size=self.sizes["formula_legend"], bold=True,
                   color=self.p["muted"], first=True, space_after=8)
        for item in (data.get("legend") or [])[:6]:
            self._para(tf, f"   {item}", size=self.sizes["formula_legend"],
                       space_after=7)
        unit = data.get("units")
        if unit:
            self._para(tf, f"Единица измерения: {unit}",
                       size=self.sizes["formula_legend"], italic=True,
                       color=self.p["muted"], space_before=8, space_after=0)
        return slide

    def slide_table(self, data):
        slide = self._blank()
        self._title(slide, data["heading"])

        rows_data = data["rows"]
        header = data.get("header")
        n_rows = len(rows_data) + (1 if header else 0)
        n_cols = len(header) if header else len(rows_data[0])

        y = self.body_top + Inches(0.15)
        avail_h = self.footer_y - y - Inches(0.45)
        gfx = slide.shapes.add_table(n_rows, n_cols, self.left, y,
                                     self.content_w, avail_h)
        table = gfx.table

        widths = data.get("widths")
        if widths:
            total = sum(widths)
            for i, wt in enumerate(widths):
                table.columns[i].width = Emu(int(self.content_w * wt / total))

        row_h = Emu(int(avail_h / n_rows))
        for r in table.rows:
            r.height = row_h

        def fill_cell(cell, text, *, bold=False, color=None, bg=None,
                      align=PP_ALIGN.LEFT, size=None):
            cell.margin_left = Inches(0.12)
            cell.margin_right = Inches(0.1)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if bg:
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(bg)
            else:
                cell.fill.background()
            tf = cell.text_frame
            tf.word_wrap = True
            para = tf.paragraphs[0]
            para.alignment = align
            run = para.add_run()
            run.text = str(text)
            run.font.size = Pt(size or self.sizes["table"])
            run.font.bold = bold
            run.font.name = self.font
            run.font.color.rgb = _rgb(color or self.p["text"])

        offset = 0
        if header:
            for c, text in enumerate(header):
                fill_cell(table.cell(0, c), text, bold=True,
                          color=self.p["white"], bg=self.p["primary"],
                          align=PP_ALIGN.CENTER)
            offset = 1

        aligns = data.get("align") or []
        for r, row in enumerate(rows_data):
            bg = self.p["surface"] if r % 2 == 0 else None
            for c, text in enumerate(row):
                al = PP_ALIGN.LEFT
                if c < len(aligns):
                    al = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                          "r": PP_ALIGN.RIGHT}.get(aligns[c], PP_ALIGN.LEFT)
                fill_cell(table.cell(r + offset, c), text, bg=bg, align=al,
                          bold=bool(data.get("bold_first_col") and c == 0))

        note = data.get("note")
        if note:
            tf = self._textbox(slide, self.left, y + avail_h + Inches(0.1),
                               self.content_w, Inches(0.35))
            self._para(tf, note, size=self.sizes["footer"] + 3,
                       color=self.p["muted"], italic=True, first=True,
                       space_after=0)
        return slide

    def slide_two_col(self, data):
        slide = self._blank()
        self._title(slide, data["heading"])

        y = self.body_top + Inches(0.15)
        gap = Inches(0.4)
        col_w = Emu(int((self.content_w - gap) / 2))
        col_h = self.footer_y - y - Inches(0.25)

        for idx, col in enumerate(data["columns"][:2]):
            x = self.left + (col_w + gap) * idx
            head_h = Inches(0.62)
            self._rect(slide, x, y, col_w, head_h, self.p["primary"])
            tf = self._textbox(slide, x + Inches(0.2), y, col_w - Inches(0.4),
                               head_h, anchor=MSO_ANCHOR.MIDDLE)
            self._para(tf, col["title"], size=self.sizes["subtitle"], bold=True,
                       color=self.p["white"], align=PP_ALIGN.CENTER, first=True,
                       space_after=0)

            body_y = y + head_h
            body_h = col_h - head_h
            self._rect(slide, x, body_y, col_w, body_h, self.p["surface"])
            tf = self._textbox(slide, x + Inches(0.22), body_y + Inches(0.22),
                               col_w - Inches(0.44), body_h - Inches(0.4))
            for i, item in enumerate((col.get("items") or [])[:5]):
                self._para(tf, f"•  {item}", size=self.sizes["body_small"],
                           first=(i == 0), space_after=11)
        return slide

    def slide_scheme(self, data):
        """Схема из фигур: цепочка блоков со стрелками или ряд блоков."""
        slide = self._blank()
        self._title(slide, data["heading"])

        blocks = data["blocks"]
        y = self.body_top + Inches(0.45)
        chain = data.get("style", "chain") == "chain"

        n = len(blocks)
        gap = Inches(0.20) if chain else Inches(0.3)
        arrow_w = Inches(0.52) if chain else Emu(0)
        total_gaps = (gap * (n - 1)) + (arrow_w * (n - 1) if chain else 0)
        box_w = Emu(int((self.content_w - total_gaps) / n))
        box_h = Inches(1.5)

        x = self.left
        for i, block in enumerate(blocks):
            accent = block.get("accent", False)
            fill = self.p["primary"] if accent else self.p["accent_soft"]
            text_color = self.p["white"] if accent else self.p["primary"]
            self._rect(slide, x, y, box_w, box_h, fill,
                       shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            tf = self._textbox(slide, x + Inches(0.13), y + Inches(0.12),
                               box_w - Inches(0.26), box_h - Inches(0.24),
                               anchor=MSO_ANCHOR.MIDDLE)
            self._para(tf, block["title"], size=self.sizes["body_small"],
                       bold=True, color=text_color, align=PP_ALIGN.CENTER,
                       first=True, space_after=0)

            sub = block.get("sub")
            if sub:
                tf2 = self._textbox(slide, x, y + box_h + Inches(0.12),
                                    box_w, Inches(0.9))
                self._para(tf2, sub, size=self.sizes["footer"] + 5,
                           color=self.p["muted"], align=PP_ALIGN.CENTER,
                           first=True, space_after=0)

            x += box_w
            if chain and i < n - 1:
                ax = x + Emu(int(gap / 2))
                self._rect(slide, ax, y + Emu(int(box_h / 2)) - Inches(0.14),
                           arrow_w, Inches(0.28), self.p["accent"],
                           shape=MSO_SHAPE.RIGHT_ARROW)
                x += arrow_w + gap
            elif i < n - 1:
                x += gap

        note = data.get("note")
        if note:
            ny = y + box_h + Inches(1.15)
            tf = self._textbox(slide, self.left, ny, self.content_w,
                               self.footer_y - ny - Inches(0.15))
            self._para(tf, note, size=self.sizes["body_small"],
                       color=self.p["muted"], italic=True, first=True,
                       space_after=0)
        return slide

    # Высоты для раскладки слайда-примера, в дюймах.
    # Строка текста размером S pt занимает примерно S × 1,22 / 72 дюйма;
    # к ней добавляется межабзацный отступ. Значения ниже посчитаны так
    # и сверены рендером scripts/preview_deck.py.
    GIVEN_LINE_IN = 0.37       # строка «Дано», 20 pt + отступ 2 pt
    GIVEN_PAD_IN = 0.62        # заголовок «Дано:» и поля блока
    GIVEN_GAP_IN = 0.20        # отступ от блока «Дано» до шагов
    STEP_LABEL_IN = 0.37       # название действия, 20 pt
    STEP_CALC_IN = 0.40        # строка расчёта, 22 pt
    STEP_COMMENT_IN = 0.36     # строка комментария, 20 pt
    STEP_GAP_IN = 0.10         # отступ между действиями
    ANSWER_H_IN = 0.62         # блок «Ответ» в одну строку
    ANSWER_LINE2_IN = 0.38     # добавка на вторую строку ответа
    ANSWER_GAP_IN = 0.14
    # предельная длина строки: «Ответ: » + текст, 22 pt на всю ширину
    ANSWER_WRAP_CHARS = 78
    COMMENT_WRAP_CHARS = 90

    @classmethod
    def example_height_in(cls, data) -> float:
        """Требуемая высота содержимого слайда-примера, в дюймах."""
        total = 0.0
        given = data.get("given") or []
        if given:
            total += (cls.GIVEN_PAD_IN + cls.GIVEN_LINE_IN * len(given)
                      + cls.GIVEN_GAP_IN)
        for i, step in enumerate(data.get("steps") or []):
            if i:
                total += cls.STEP_GAP_IN
            total += cls.STEP_LABEL_IN
            if step.get("calc"):
                total += cls.STEP_CALC_IN
            comment = step.get("comment")
            if comment:
                lines = 1 + len(str(comment)) // cls.COMMENT_WRAP_CHARS
                total += cls.STEP_COMMENT_IN * lines
        if data.get("answer"):
            total += cls.ANSWER_H_IN + cls.ANSWER_GAP_IN
            if len(str(data["answer"])) > cls.ANSWER_WRAP_CHARS:
                total += cls.ANSWER_LINE2_IN
        return total

    @classmethod
    def example_available_in(cls, theme) -> float:
        """Свободная высота под содержимое примера, в дюймах."""
        m = theme["pptx"]["margins_in"]
        title_block = 1.05 + 0.28          # заголовок и акцентная линия
        return (theme["pptx"]["slide_height_in"] - m["top"] - title_block
                - m["bottom"] - 0.10)

    def slide_example(self, data):
        """Слайд разобранного числового примера.

        Блок «Ответ» прижат к низу, поэтому под шаги отводится только
        оставшаяся высота — иначе текст шагов заходил бы на ответ.
        """
        slide = self._blank()
        self._title(slide, data["heading"])

        answer = data.get("answer")
        answer_h = Inches(self.ANSWER_H_IN) if answer else Emu(0)
        if len(str(answer or "")) > self.ANSWER_WRAP_CHARS:
            answer_h += Inches(self.ANSWER_LINE2_IN)
        # нижняя граница области шагов
        steps_bottom = self.footer_y - Inches(0.10)
        if answer:
            steps_bottom -= answer_h + Inches(self.ANSWER_GAP_IN)

        y = self.body_top + Inches(0.02)
        given = data.get("given")
        if given:
            gh = (Inches(self.GIVEN_PAD_IN)
                  + Inches(self.GIVEN_LINE_IN) * len(given))
            self._rect(slide, self.left, y, self.content_w, gh,
                       self.p["surface"])
            tf = self._textbox(slide, self.left + Inches(0.25), y + Inches(0.11),
                               self.content_w - Inches(0.5), gh - Inches(0.22))
            self._para(tf, "Дано (данные условные):",
                       size=self.sizes["body_small"], bold=True,
                       color=self.p["muted"], first=True, space_after=4)
            for g in given:
                self._para(tf, f"   {g}", size=self.sizes["body_small"],
                           space_after=2)
            y += gh + Inches(self.GIVEN_GAP_IN)

        steps = data.get("steps") or []
        tf = self._textbox(slide, self.left, y, self.content_w,
                           max(steps_bottom - y, Inches(0.4)))
        first = True
        for i, step in enumerate(steps, 1):
            label = step.get("label", f"Действие {i}")
            self._para(tf, f"{i}. {label}", size=self.sizes["body_small"],
                       bold=True, color=self.p["primary"], first=first,
                       space_after=2, space_before=0 if first else 7)
            first = False
            calc = step.get("calc")
            if calc:
                self._para(tf, f"     {calc}", size=self.sizes["body"],
                           bold=True, font=self.cfg["font_mono"], space_after=2)
            comment = step.get("comment")
            if comment:
                self._para(tf, f"     {comment}",
                           size=self.sizes["body_small"],
                           color=self.p["muted"], italic=True, space_after=1)

        if answer:
            ay = self.footer_y - answer_h - Inches(0.10)
            self._rect(slide, self.left, ay, self.content_w, answer_h,
                       self.p["accent_soft"])
            self._rect(slide, self.left, ay, Emu(int(Inches(0.07))), answer_h,
                       self.p["primary"])
            tfa = self._textbox(slide, self.left + Inches(0.3), ay,
                                self.content_w - Inches(0.6), answer_h,
                                anchor=MSO_ANCHOR.MIDDLE)
            self._para(tfa, f"Ответ: {answer}", size=self.sizes["body"],
                       bold=True, color=self.p["primary"], first=True,
                       space_after=0)
        return slide

    def slide_check(self, data):
        slide = self._blank()
        self._title(slide, data.get("heading", "Закрепление"))
        y = self.body_top + Inches(0.15)
        tf = self._textbox(slide, self.left, y, self.content_w,
                           self.footer_y - y - Inches(0.2))
        first = True
        for i, q in enumerate((data.get("questions") or [])[:5], 1):
            self._para(tf, f"{i}.  {q}", size=self.sizes["body"], first=first,
                       space_after=16)
            first = False
        return slide

    def slide_summary(self, data):
        slide = self._blank()
        self._title(slide, data.get("heading", "Итог занятия"))
        y = self.body_top + Inches(0.15)
        items = (data.get("bullets") or [])[:5]
        tf = self._textbox(slide, self.left, y, self.content_w,
                           self.footer_y - y - Inches(0.2))
        first = True
        for item in items:
            # тире, а не «галочка»: символ ✓ есть не во всех версиях Calibri
            self._para(tf, f"—  {item}", size=self.sizes["body"], first=first,
                       space_after=15)
            first = False
        return slide

    def slide_homework(self, data):
        slide = self._blank()
        self._title(slide, data.get("heading", "Домашнее задание"))
        y = self.body_top + Inches(0.2)
        h = Inches(2.2)
        self._rect(slide, self.left, y, self.content_w, h, self.p["surface"])
        self._rect(slide, self.left, y, Emu(int(Inches(0.07))), h,
                   self.p["accent"])
        tf = self._textbox(slide, self.left + Inches(0.32), y + Inches(0.26),
                           self.content_w - Inches(0.64), h - Inches(0.5))
        self._para(tf, data["text"], size=self.sizes["body"], first=True,
                   space_after=10)
        for item in (data.get("bullets") or [])[:4]:
            self._para(tf, f"•  {item}", size=self.sizes["body_small"],
                       space_after=8)
        source = data.get("source")
        if source:
            y2 = y + h + Inches(0.35)
            tf = self._textbox(slide, self.left, y2, self.content_w,
                               self.footer_y - y2 - Inches(0.15))
            self._para(tf, source, size=self.sizes["body_small"],
                       color=self.p["muted"], italic=True, first=True,
                       space_after=0)
        return slide

    # --- сборка -------------------------------------------------------

    _DISPATCH = {
        "title": "slide_title",
        "question": "slide_question",
        "bullets": "slide_bullets",
        "definition": "slide_definition",
        "formula": "slide_formula",
        "table": "slide_table",
        "two_col": "slide_two_col",
        "scheme": "slide_scheme",
        "example": "slide_example",
        "check": "slide_check",
        "summary": "slide_summary",
        "homework": "slide_homework",
    }

    def build(self, slides: list[dict], out_path):
        for index, data in enumerate(slides, 1):
            kind = data.get("type")
            method = self._DISPATCH.get(kind)
            if method is None:
                raise ValueError(
                    f"Занятие {self.number}: неизвестный тип слайда «{kind}»")
            slide = getattr(self, method)(data)
            if kind != "title":
                self._footer(slide, index)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(out_path))
        return out_path


def build_deck(theme, discipline, number, topic, slides, out_path):
    return DeckBuilder(theme, discipline, number, topic).build(slides, out_path)
