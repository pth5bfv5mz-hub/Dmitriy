"""Рендер слайдов собранной презентации в изображения.

Приблизительный рендер по геометрии фигур из .pptx: нужен, чтобы
контролировать вёрстку и читать слайды без PowerPoint. Шрифты подставляются
системные, поэтому переносы строк могут отличаться от PowerPoint на
несколько пикселей.

Используется двумя скриптами: preview_deck.py (сохраняет PNG) и
build_mobile.py (встраивает изображения в страницу для телефона).
"""
from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

FONT_REG = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
FONT_MONO = "/usr/share/fonts/truetype/liberation/LiberationMono-Bold.ttf"

_font_cache: dict = {}


def _font(size_pt: float, scale: int, *, bold=False, mono=False):
    px = max(7, int(size_pt * scale / 72))
    key = (px, bold, mono)
    if key not in _font_cache:
        path = FONT_MONO if mono else (FONT_BOLD if bold else FONT_REG)
        _font_cache[key] = ImageFont.truetype(path, px)
    return _font_cache[key]


def _px(emu, scale: int) -> int:
    return int(Emu(int(emu)).inches * scale)


def _wrap(draw, text, fnt, max_w):
    words, lines, cur = str(text).split(), [], ""
    for w in words:
        probe = f"{cur} {w}".strip()
        if draw.textlength(probe, font=fnt) <= max_w or not cur:
            cur = probe
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def _shape_color(obj, default=None):
    try:
        if obj.fill.type == 1:
            return "#" + str(obj.fill.fore_color.rgb)
    except Exception:
        pass
    return default


def _run_color(run, default="#1A1A1A"):
    try:
        if run.font.color and run.font.color.rgb:
            return "#" + str(run.font.color.rgb)
    except Exception:
        pass
    return default


def _draw_table(draw, shape, scale, warn, slide_no):
    table = shape.table
    x0, y0 = _px(shape.left, scale), _px(shape.top, scale)
    col_w = [_px(c.width, scale) for c in table.columns]
    row_h = [_px(r.height, scale) for r in table.rows]
    y = y0
    for r, row in enumerate(table.rows):
        x = x0
        for c, cell in enumerate(row.cells):
            w, h = col_w[c], row_h[r]
            draw.rectangle([x, y, x + w, y + h], fill=_shape_color(cell),
                           outline="#B9C4D4")
            runs = [run for p in cell.text_frame.paragraphs for run in p.runs]
            if runs:
                run = runs[0]
                size = run.font.size.pt if run.font.size else 18
                fnt = _font(size, scale, bold=bool(run.font.bold))
                lines = _wrap(draw, cell.text, fnt, w - 14)
                lh = int(size * scale / 72 * 1.22)
                ty = y + max(3, (h - lh * len(lines)) // 2)
                for line in lines:
                    draw.text((x + 7, ty), line, font=fnt,
                              fill=_run_color(run))
                    ty += lh
                if lh * len(lines) > h:
                    warn.append(f"слайд {slide_no}: текст ячейки таблицы "
                                f"не вмещается по высоте: «{cell.text[:40]}»")
            x += w
        y += row_h[r]


def _draw_textframe(draw, shape, scale, warn, slide_no, canvas_h):
    x, y = _px(shape.left, scale), _px(shape.top, scale)
    w = _px(shape.width, scale)
    cursor = y
    for para in shape.text_frame.paragraphs:
        runs = para.runs
        if not runs:
            continue
        run = runs[0]
        size = run.font.size.pt if run.font.size else 18
        mono = bool(run.font.name and "Consolas" in str(run.font.name))
        fnt = _font(size, scale, bold=bool(run.font.bold), mono=mono)
        text = "".join(r.text for r in runs)
        cursor += int((para.space_before.pt if para.space_before else 0)
                      * scale / 72)
        lh = int(size * scale / 72 * 1.22)
        align = str(para.alignment or "")
        for line in _wrap(draw, text, fnt, w - 4):
            tw = draw.textlength(line, font=fnt)
            tx = x
            if "CENTER" in align:
                tx = x + (w - tw) / 2
            elif "RIGHT" in align:
                tx = x + w - tw
            draw.text((tx, cursor), line, font=fnt, fill=_run_color(run))
            cursor += lh
        cursor += int((para.space_after.pt if para.space_after else 0)
                      * scale / 72)
    if cursor > canvas_h - 4:
        warn.append(f"слайд {slide_no}: текст выходит за нижний край слайда "
                    f"({cursor - canvas_h} px)")


def render_slides(pptx_path: Path, scale: int = 110):
    """Возвращает (список изображений PIL, список предупреждений)."""
    prs = Presentation(str(pptx_path))
    W, H = _px(prs.slide_width, scale), _px(prs.slide_height, scale)
    warn: list[str] = []
    images = []

    for i, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (W, H), "white")
        draw = ImageDraw.Draw(img)
        draw.rectangle([0, 0, W - 1, H - 1], outline="#D0D7E2")

        for shape in slide.shapes:
            if shape.has_table:
                _draw_table(draw, shape, scale, warn, i)
                continue
            if shape.shape_type is not None and not shape.has_text_frame:
                continue
            # 17 — текстовое поле; остальные автофигуры рисуем с заливкой
            if shape.shape_type is not None and shape.shape_type != 17:
                fill = _shape_color(shape)
                if fill:
                    draw.rectangle([
                        _px(shape.left, scale), _px(shape.top, scale),
                        _px(shape.left, scale) + _px(shape.width, scale),
                        _px(shape.top, scale) + _px(shape.height, scale),
                    ], fill=fill)
            if shape.has_text_frame and shape.text_frame.text.strip():
                _draw_textframe(draw, shape, scale, warn, i, H)
            right = _px(shape.left, scale) + _px(shape.width, scale)
            bottom = _px(shape.top, scale) + _px(shape.height, scale)
            if right > W + 2 or bottom > H + 2:
                warn.append(f"слайд {i}: фигура выходит за границы слайда "
                            f"(право {right - W} px, низ {bottom - H} px)")
        images.append(img)

    return images, list(dict.fromkeys(warn))


def contact_sheet(images, cols: int = 3):
    """Все слайды одной картинкой — для быстрой вычитки."""
    if not images:
        return None
    W, H = images[0].size
    tw, th = W // 2, H // 2
    rows = (len(images) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * tw, rows * th), "#E9EDF3")
    for idx, im in enumerate(images):
        sheet.paste(im.resize((tw, th)), ((idx % cols) * tw, (idx // cols) * th))
    return sheet
